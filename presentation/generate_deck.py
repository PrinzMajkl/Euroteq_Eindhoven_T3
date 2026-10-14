import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor

OUTPUT_PATH = os.path.join(os.path.dirname(__file__), "plumetrace_presentation.pptx")
LOGO_PATH = os.path.join(os.path.dirname(__file__), "assets", "logo.jpeg")
ICONS_DIR = os.path.join(os.path.dirname(__file__), "assets", "icons")

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank_layout = prs.slide_layouts[6]

# Brand Palette
NAVY = RGBColor(2, 48, 71)          # #023047
CYAN = RGBColor(33, 158, 188)        # #219ebc
SKY = RGBColor(142, 202, 230)        # #8ecae6
GOLD = RGBColor(255, 183, 3)         # #ffb703
ORANGE = RGBColor(251, 133, 0)       # #fb8500
WHITE = RGBColor(255, 255, 255)      # #ffffff
CARD_BG = RGBColor(244, 249, 252)    # #f4f9fc soft tint
BORDER_CLR = RGBColor(180, 218, 238) # clean subtle border

FONT_NAME = "Calibri"

def set_background(slide):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = WHITE

def add_slide_header(slide, title_text):
    tb = slide.shapes.add_textbox(Inches(0.8), Inches(0.55), Inches(11.733), Inches(0.85))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.TOP
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.text = title_text
    p.alignment = PP_ALIGN.LEFT
    p.font.name = FONT_NAME
    p.font.size = Pt(38)
    p.font.bold = True
    p.font.color.rgb = NAVY

def add_clean_card(slide, left, top, width, height, title, bullets, icon_name=None, subtitle=None, title_size=22, bullet_size=16):
    # Main Card Shape
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.adjustments[0] = 0.035
    shape.fill.solid()
    shape.fill.fore_color.rgb = CARD_BG
    shape.line.color.rgb = BORDER_CLR
    shape.line.width = Pt(1.5)
    
    # Optional Icon
    has_icon = False
    if icon_name:
        icon_path = os.path.join(ICONS_DIR, f"{icon_name}.png")
        if os.path.exists(icon_path):
            slide.shapes.add_picture(icon_path, left + Inches(0.32), top + Inches(0.32), width=Inches(0.52), height=Inches(0.52))
            has_icon = True

    # Title Box
    title_left = left + Inches(0.96 if has_icon else 0.35)
    title_width = width - Inches(1.25 if has_icon else 0.7)
    
    tb_title = slide.shapes.add_textbox(title_left, top + Inches(0.28), title_width, Inches(0.65))
    tf_t = tb_title.text_frame
    tf_t.word_wrap = True
    tf_t.vertical_anchor = MSO_ANCHOR.TOP
    tf_t.margin_left = tf_t.margin_top = tf_t.margin_right = tf_t.margin_bottom = 0
    
    p_main = tf_t.paragraphs[0]
    p_main.text = title
    p_main.alignment = PP_ALIGN.LEFT
    p_main.font.name = FONT_NAME
    p_main.font.size = Pt(title_size)
    p_main.font.bold = True
    p_main.font.color.rgb = NAVY
    
    if subtitle:
        p_sub = tf_t.add_paragraph()
        p_sub.text = subtitle
        p_sub.alignment = PP_ALIGN.LEFT
        p_sub.font.name = FONT_NAME
        p_sub.font.size = Pt(13)
        p_sub.font.bold = True
        p_sub.font.color.rgb = CYAN
        p_sub.space_before = Pt(3)

    # Bullets Container
    bullets_top = top + Inches(1.05 if subtitle else 0.95)
    bullets_height = height - Inches(1.15 if subtitle else 1.05)
    
    tb_b = slide.shapes.add_textbox(left + Inches(0.35), bullets_top, width - Inches(0.7), bullets_height)
    tf_b = tb_b.text_frame
    tf_b.word_wrap = True
    tf_b.vertical_anchor = MSO_ANCHOR.TOP
    tf_b.margin_left = tf_b.margin_top = tf_b.margin_right = tf_b.margin_bottom = 0
    
    for i, b in enumerate(bullets):
        bp = tf_b.paragraphs[0] if i == 0 else tf_b.add_paragraph()
        bp.text = f"-  {b}"
        bp.alignment = PP_ALIGN.LEFT
        bp.font.name = FONT_NAME
        bp.font.size = Pt(bullet_size)
        bp.font.color.rgb = NAVY
        bp.space_after = Pt(10)
        bp.space_before = Pt(0)
        
    return shape

def add_notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


# ==========================================================
# SLIDE 1: Title Slide (Section 1)
# ==========================================================
s1 = prs.slides.add_slide(blank_layout)
set_background(s1)

if os.path.exists(LOGO_PATH):
    s1.shapes.add_picture(LOGO_PATH, Inches(1.4), Inches(2.0), width=Inches(3.4), height=Inches(3.4))

s1_tb = s1.shapes.add_textbox(Inches(5.2), Inches(2.2), Inches(7.3), Inches(3.0))
tf1 = s1_tb.text_frame
tf1.word_wrap = True
tf1.vertical_anchor = MSO_ANCHOR.MIDDLE
tf1.margin_left = tf1.margin_top = tf1.margin_right = tf1.margin_bottom = 0

p1 = tf1.paragraphs[0]
p1.text = "PlumeTrace"
p1.alignment = PP_ALIGN.LEFT
p1.font.name = FONT_NAME
p1.font.size = Pt(64)
p1.font.bold = True
p1.font.color.rgb = CYAN
p1.space_after = Pt(14)

p2 = tf1.add_paragraph()
p2.text = "Autonomous Water Protection for the Pegnitz River"
p2.alignment = PP_ALIGN.LEFT
p2.font.name = FONT_NAME
p2.font.size = Pt(26)
p2.font.bold = True
p2.font.color.rgb = NAVY

add_notes(s1,
    "Good morning members of the jury and mentors. Today, we present PlumeTrace - an autonomous, "
    "closed-loop Digital Twin system designed to protect urban rivers.\n\n"
    "Across Europe, urban rivers are essential for recreation, drinking water, and ecology. "
    "Yet beneath the surface, they face sudden, unmonitored pollution events that traditional testing completely misses.\n\n"
    "Using the Pegnitz River in Nuremberg as our testbed, we demonstrate how coupling an autonomous "
    "robotic agent with a real-time predictive Digital Twin transforms water quality monitoring from "
    "blind, infrequent sampling into actionable protection."
)


# ==========================================================
# SLIDE 2: Why Pegnitz (Section 2 - 1/3)
# ==========================================================
s2 = prs.slides.add_slide(blank_layout)
set_background(s2)
add_slide_header(s2, "Why Pegnitz")

s2_w = Inches(5.666)
s2_h = Inches(2.45)
s2_top1 = Inches(1.65)
s2_top2 = Inches(4.35)
s2_left1 = Inches(0.8)
s2_left2 = Inches(6.866)

add_clean_card(s2, s2_left1, s2_top1, s2_w, s2_h, "Public Health", [
    "Raw sewage spikes after convective rainfall",
    "Instant health risk for surfers, waders, and park visitors"
], icon_name="health", title_size=21, bullet_size=16)

add_clean_card(s2, s2_left2, s2_top1, s2_w, s2_h, "Sewer Overflows (CSOs)", [
    "1,500 km municipal network overloaded by storms",
    "1.5M - 5M+ EUR cost per retention basin"
], icon_name="storm_pipe", title_size=21, bullet_size=16)

add_clean_card(s2, s2_left1, s2_top2, s2_w, s2_h, "Knoblauchsland Agriculture", [
    "Nitrate leaching into drinking water aquifers",
    "Strict EU limits and compliance penalties"
], icon_name="agriculture", title_size=21, bullet_size=16)

add_clean_card(s2, s2_left2, s2_top2, s2_w, s2_h, "Summer Heat and Drought", [
    "Low flow concentrates urban pollutants",
    "Dissolved oxygen drops < 3 mg/L (fish kills)"
], icon_name="drought", title_size=21, bullet_size=16)

add_notes(s2,
    "Why the Pegnitz River? Because it represents four urgent pressures colliding in one urban corridor:\n\n"
    "First, Public Health: Surfers at the Fuchslochwelle wave and waders face raw sewage contamination after storms, "
    "creating immediate health risks.\n"
    "Second, Urban Infrastructure: Nuremberg's 1,500 km sewer network overspills during heavy rains. "
    "Modernizing a single retention basin costs between 1.5 and 5 million euros - cities cannot afford to invest blindly.\n"
    "Third, Agriculture: Intensive vegetable farming in Knoblauchsland leaches nitrates into alluvial drinking water aquifers under strict EU penalties.\n"
    "And fourth, Climate Drought: Summer heat reduces river discharge, collapsing oxygen levels below 3 mg/L and killing fish."
)


# ==========================================================
# SLIDE 3: Current Monitoring (Section 2 - 2/3)
# (THE ONLY SLIDE WITH THE BOTTOM HIGHLIGHT BANNER)
# ==========================================================
s3 = prs.slides.add_slide(blank_layout)
set_background(s3)
add_slide_header(s3, "Current Monitoring")

col_w = Inches(5.666)
col3_h = Inches(4.0)
top_pos = Inches(1.65)

add_clean_card(s3, Inches(0.8), top_pos, col_w, col3_h, "Municipal Authorities", [
    "Temperature, pH, conductivity, turbidity",
    "Dissolved oxygen and nutrients (Nitrate, Phosphate)",
    "Ammonium (NH4+): primary raw sewage tracer"
], icon_name="municipal_sensor", subtitle="Stadtentwässerung Nürnberg (SUN) & WWA", title_size=22, bullet_size=16)

add_clean_card(s3, Inches(6.866), top_pos, col_w, col3_h, "The Sampling Gap", [
    "Automatic sensors: continuous, but confined to few fixed locations",
    "Manual grab sampling: often only 1x per month (24-48h lab delay)",
    "Convective storm overflows: last only 2 to 4 hours"
], icon_name="sampling_gap", title_size=22, bullet_size=16)

# The Result (Left aligned, subtle 0.035 radius)
res_box = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(6.0), Inches(11.733), Inches(0.75))
res_box.adjustments[0] = 0.035
res_box.fill.solid()
res_box.fill.fore_color.rgb = CARD_BG
res_box.line.color.rgb = CYAN
res_box.line.width = Pt(1.5)

tf_res = res_box.text_frame
tf_res.word_wrap = True
tf_res.vertical_anchor = MSO_ANCHOR.MIDDLE
tf_res.margin_left = Inches(0.35)
tf_res.margin_right = Inches(0.35)
tf_res.margin_top = Inches(0.1)

p_res = tf_res.paragraphs[0]
p_res.alignment = PP_ALIGN.LEFT
p_res.text = "The Result:  "
p_res.font.name = FONT_NAME
p_res.font.size = Pt(18)
p_res.font.bold = True
p_res.font.color.rgb = ORANGE

run_res = p_res.add_run()
run_res.text = "Plume arrives, causes harm, and clears undetected."
run_res.font.name = FONT_NAME
run_res.font.size = Pt(18)
run_res.font.bold = True
run_res.font.color.rgb = NAVY

add_notes(s3,
    "To address this, we looked at what municipal authorities currently measure.\n\n"
    "In Nuremberg, the drainage service SUN tracks eight standard parameters: temperature, pH, conductivity, turbidity, "
    "dissolved oxygen, and nutrients like ammonium, nitrate, and phosphate. Ammonium is the critical chemical fingerprint of raw sewage.\n\n"
    "The sampling gap is two-fold: automated sensors are confined to a few fixed points, while manual sampling happens as rarely as once a month with days of lab delay. "
    "The result: acute plumes arrive, cause damage, and clear completely undetected."
)


# ==========================================================
# SLIDE 4: Key Stakeholders (Section 2 - 3/3)
# ==========================================================
s4 = prs.slides.add_slide(blank_layout)
set_background(s4)
add_slide_header(s4, "Key Stakeholders")

s4_card_h = Inches(4.8)
add_clean_card(s4, Inches(0.8), Inches(1.65), col_w, s4_card_h, "Water Authorities (WWA, SUN)", [
    "Cannot locate pollution origin (fixed stations lack spatial tracking)",
    "Blind infrastructure spending (1.5M - 5M+ EUR per basin without plume data)",
    "Chronic pipe seepage undetected until structural failure",
    "Lab results arrive 24-48 hours too late for emergency action"
], icon_name="authority", title_size=22, bullet_size=16)

add_clean_card(s4, Inches(6.866), Inches(1.65), col_w, s4_card_h, "Citizens and Surfers", [
    "Zero real-time visibility into water safety",
    "Unknowingly enter water during toxic overflow events",
    "High risk of gastrointestinal and skin infections",
    "Delayed municipal reports provide no immediate guidance"
], icon_name="citizens_surfer", title_size=22, bullet_size=16)

add_notes(s4,
    "This monitoring failure creates acute problems for our two key stakeholder groups today:\n\n"
    "First, Water Authorities like SUN and WWA: They must invest millions of euros into retention basins, "
    "but without spatial plume tracking, they are forced to invest blindly without knowing which overflow pipes cause the greatest damage. "
    "Furthermore, laboratory delays mean they only learn of acute spills days after they have already washed away.\n\n"
    "Second, Citizens and Surfers: People entering the Pegnitz at the Fuchslochwelle or in city parks have zero real-time warning. "
    "They unknowingly expose themselves to acute bacterial contamination because municipal data is either delayed by days or buried in complex chemical reports."
)


# ==========================================================
# SLIDE 5: The PlumeTrace Architecture (Section 3 - 1/3)
# (FULL OVERARCHING SOLUTION: ROBOT IN RIVER + TWIN + DASHBOARD)
# ==========================================================
s5 = prs.slides.add_slide(blank_layout)
set_background(s5)
add_slide_header(s5, "The PlumeTrace Architecture")

col3_w = Inches(3.644)
col3_gap = Inches(0.4)
s5_card_h = Inches(4.8)
s5_top = Inches(1.65)

add_clean_card(s5, Inches(0.8), s5_top, col3_w, s5_card_h, "Autonomous River Robot", [
    "Autonomous Surface Vessel (ASV) on Pegnitz",
    "Continuous chemical payload (Ammonium, DO, pH)",
    "360° LiDAR scans riverbank morphology & culverts",
    "Edge autonomy with fail-safe return-to-dock"
], icon_name="river_asv", title_size=20, bullet_size=15)

add_clean_card(s5, Inches(0.8) + col3_w + col3_gap, s5_top, col3_w, s5_card_h, "Predictive Digital Twin", [
    "Hydrodynamic river advection-dispersion model",
    "Ingests live DWD storm radar & river discharge",
    "Multi-objective solver computes optimal sampling",
    "Continuously directs robot where to navigate next"
], icon_name="cloud_twin", title_size=20, bullet_size=15)

add_clean_card(s5, Inches(0.8) + (col3_w + col3_gap) * 2, s5_top, col3_w, s5_card_h, "Municipal Command Center", [
    "Geospatial 3D twin with Bavarian LOD2 data",
    "Real-time pollution heatmaps & outfall health",
    "Pinpoints exact sewer outfalls causing violations",
    "Automated statutory alerts (<60s) to WWA & SUN"
], icon_name="digital_twin_3d", title_size=20, bullet_size=15)

add_notes(s5,
    "Here is our overarching technical solution: PlumeTrace connects physical robotics, cloud digital twinning, and municipal decision-making into one unified ecosystem.\n\n"
    "First, on the water: An Autonomous Surface Vessel continuously patrols the Pegnitz corridor, equipped with chemical sensors for ammonium and dissolved oxygen, and 360-degree LiDAR mapping riverbank morphology.\n\n"
    "Second, in the cloud: A predictive Digital Twin runs hydrodynamic advection-dispersion modeling. By assimilating real-time weather radar from the German Weather Service (DWD) and river discharge, the twin anticipates where storm runoff will arrive and tells the robot where to navigate next.\n\n"
    "Third, in the city hall: A 3D geospatial dashboard built on CesiumJS and Bavarian LOD2 OpenData translates raw data into outfall asset health scores, enabling targeted infrastructure investment and immediate public warnings."
)


# ==========================================================
# SLIDE 6: Autonomous Closed-Loop Guidance (Section 3 - 2/3)
# ==========================================================
s6 = prs.slides.add_slide(blank_layout)
set_background(s6)
add_slide_header(s6, "Autonomous Closed-Loop Guidance")

add_clean_card(s6, Inches(0.8), Inches(1.65), col_w, s4_card_h, "4D Adaptive Retasking", [
    "Dynamically recomputes waypoints in real time",
    "Chemical Gradient: homes upstream to leaking pipe",
    "Storm Rain Radar: pre-positions vessel before runoff",
    "Data Freshness: prioritizes unvisited river reaches",
    "Dynamic goal updates dispatched in under 35 seconds"
], icon_name="adaptive_retask", title_size=22, bullet_size=16)

add_clean_card(s6, Inches(6.866), Inches(1.65), col_w, s4_card_h, "Why Digital Twinning?", [
    "Active Gradient Tracing overcomes point dilution",
    "Predictive storm anticipation replaces reactive lab lag",
    "One mobile vessel replaces 50+ biofouled static probes",
    "Continuous feedback synchronizes twin with physical state",
    "Empirical evidence justifies targeted retention investments"
], icon_name="safety_resilience", title_size=22, bullet_size=16)

add_notes(s6,
    "The core innovation of PlumeTrace is its Autonomous Closed Loop.\n\n"
    "Instead of driving a static patrol line, our system continuously recomputes waypoints based on a 4D objective function: "
    "tracking chemical gradients upstream, anticipating rain runoff via live weather radar, and ensuring older areas are revisited.\n\n"
    "Why does digital twinning beat traditional static sensors? Static probes suffer from point dilution: they tell you pollution passed, but cannot identify where it originated. "
    "By coupling physical sensing with predictive twin simulation, we pinpoint the exact leaking outfall in minutes."
)


# ==========================================================
# SLIDE 7: Stakeholder Action & AI Public Access (Section 3 - 3/3)
# ==========================================================
s7 = prs.slides.add_slide(blank_layout)
set_background(s7)
add_slide_header(s7, "Stakeholder Action & AI Public Access")

add_clean_card(s7, Inches(0.8), Inches(1.65), col_w, s4_card_h, "City & Authority Command", [
    "Outfall Asset Health: links chemical spikes to specific pipe IDs",
    "Targeted CapEx: empirical proof for 1.5M - 5M+ EUR basin upgrades",
    "Chronic Seepage: uncovers leaking pipes before structural collapse",
    "Automated compliance logs for EU Water Framework Directive & OGewV"
], icon_name="authority", title_size=22, bullet_size=16)

add_clean_card(s7, Inches(6.866), Inches(1.65), col_w, s4_card_h, "AI Public Access via MCP", [
    "Model Context Protocol (MCP) connects AI models to live twin data",
    "Conversational Queries: 'Can I surf at Fuchslochwelle today?'",
    "Live Citizen Traffic Light: instant Green / Yellow / Red alerts",
    "Eliminates 48h lab delay with direct public health warnings"
], icon_name="ai_mcp", title_size=22, bullet_size=16)

add_notes(s7,
    "Finally, how do stakeholders interact with PlumeTrace?\n\n"
    "For Water Authorities (SUN & WWA): We provide concrete asset intelligence. By assigning health scores to specific outfall IDs, "
    "we provide the empirical justification needed to prioritize multi-million euro retention basin expansions.\n\n"
    "For Citizens and Surfers: We expose the Digital Twin via the Model Context Protocol (MCP). Anyone can ask an AI assistant: "
    "'Can I surf at the Fuchslochwelle wave today?' or 'Is the river safe for my dog right now?'. "
    "The AI queries live telemetry and forecasts, returning a clear traffic light advisory (Green, Yellow, or Red). "
    "This transforms complex environmental chemistry into immediate, transparent public safety."
)

prs.save(OUTPUT_PATH)
print(f"Successfully generated clean PowerPoint presentation at: {OUTPUT_PATH}")
